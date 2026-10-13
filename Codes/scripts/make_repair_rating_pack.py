"""Blinded teacher-rating pack for the repair arms (reviewer Q1; optional, needs new annotation).

24 matched pairs of R2-T and R5 replies (same item, tutor and language), 12 Hindi and 12 Bengali,
2 per tutor and language, drawn with a fixed seed before any rating exists. Each reply is rated on
the five PQI-R dimensions plus the naturalness of its language, and each pair gets a blinded
preference. The pack holds no arm, tutor or study labels; the key is written separately and must
not be sent.

    PYTHONPATH=src python scripts/make_repair_rating_pack.py
Writes review/outgoing/repair_rating_pack.xlsx and review/repair_rating_key.csv.
"""
import json
import random

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from tlt.config import load_run
from tlt.io import read_jsonl, read_table

run = load_run("full")
rng = random.Random(run.seed + 71)
rep = read_table(run.path("results", "repair_judgements_aggregated.csv"), low_memory=False)
rep = rep[(rep.variant == "original") & (rep.cue == "C0") & rep.condition.isin(["HI", "BN"])
          & rep.arm.isin(["R2T", "R5"])]
pairs = rep.pivot_table(index=["item_id", "condition", "model"], columns="arm", values="target_id",
                        aggfunc="first").dropna().reset_index()
chosen = []
for (lang, model), g in pairs.sort_values(["condition", "model", "item_id"]).groupby(["condition", "model"]):
    chosen += rng.sample(list(g.itertuples(index=False)), 2)
texts = {}
for line in open(run.path("runs", "repair_generations.jsonl"), encoding="utf-8"):
    r = json.loads(line)
    texts[r["call_id"]] = r
items = {it["item_id"]: it for it in read_jsonl(run.path("data", "items_multilingual.jsonl"))}
prompts = {p["prompt_id"]: p for p in read_jsonl(run.path("data", "prompts.jsonl"))}

rows, key = [], []
rng.shuffle(chosen)
for n, p in enumerate(chosen, start=1):
    order = ["R2T", "R5"]
    rng.shuffle(order)
    msg = prompts[f"A|{p.item_id}|{p.condition}|C0"]["user"]
    v = items[p.item_id]["versions"][p.condition]
    for label, arm in zip("AB", order):
        rid = f"P{n:02d}{label}"
        rows.append({"row_id": rid, "pair": f"P{n:02d}", "reply": label,
                     "language": {"HI": "Hindi", "BN": "Bengali"}[p.condition],
                     "correct_answer": v["correct_answer"], "student_message": msg,
                     "tutor_reply": texts[getattr(p, arm)].get("response_text") or ""})
        key.append({"row_id": rid, "pair": f"P{n:02d}", "arm": arm, "item_id": p.item_id, "language": p.condition,
                    "model": p.model, "target_id": getattr(p, arm)})
pd.DataFrame(key).to_csv(run.path("review", "repair_rating_key.csv"), index=False)

LAB = {"Mistake_Identification": "Yes,To some extent,No", "Mistake_Location": "Yes,To some extent,No",
       "Revealing_of_the_Answer": "No,Yes (correct),Yes (incorrect)",
       "Providing_Guidance": "Yes,To some extent,No", "Math_Correctness": "Yes,No",
       "Language_Naturalness": "Natural,Understandable but unnatural,Hard to understand"}
wb = Workbook()
ins = wb.active
ins.title = "Instructions"
for i, line in enumerate([
    "Rating tutor replies to a student's maths mistake",
    "Each row is one tutor reply to a student who sent the message shown, with a wrong answer.",
    "The correct answer is given for your reference; the tutor did not see it.",
    "Rate every reply on the six columns (use the drop-down lists), using the same definitions as the",
    "earlier rating round: Mistake identification: does the reply recognise that the answer is wrong?",
    "Mistake location: does it point to where the reasoning went wrong? Revealing the answer: does it",
    "state the correct (or an incorrect) final answer? Providing guidance: does it help the student",
    "work towards the answer? Mathematical correctness: is everything the tutor states mathematically",
    "correct? Language naturalness: does the reply read like natural Hindi/Bengali to a student?",
    "In the sheet 'Pairs', say which of the two replies (A or B) to the same message would help the",
    "student more, or 'No difference'. Rate the replies independently first, then the pairs.",
    "Please rate only replies in a language you read fluently. Do not discuss ratings with other raters."], 1):
    ins.cell(row=i, column=1, value=line).font = Font(bold=(i == 1))
ws = wb.create_sheet("Ratings")
cols = ["row_id", "pair", "reply", "language", "correct_answer", "student_message", "tutor_reply"] + list(LAB)
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
for r in rows:
    ws.append([r[c] for c in cols[:7]] + [""] * len(LAB))
for j, (dim, opts) in enumerate(LAB.items(), start=8):
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    ws.add_data_validation(dv)
    col = ws.cell(row=1, column=j).column_letter
    dv.add(f"{col}2:{col}{len(rows) + 1}")
for col, w in zip("ABCDEFGHIJKLM", (8, 6, 6, 9, 16, 50, 80, 14, 14, 16, 14, 14, 18)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ps = wb.create_sheet("Pairs")
ps.append(["pair", "language", "Which reply would help the student more?"])
for c in ps[1]:
    c.font = Font(bold=True)
for n in sorted({r["pair"] for r in rows}):
    ps.append([n, next(r["language"] for r in rows if r["pair"] == n), ""])
dv = DataValidation(type="list", formula1='"A,B,No difference"', allow_blank=True)
ps.add_data_validation(dv)
dv.add(f"C2:C{len({r['pair'] for r in rows}) + 1}")
ps.column_dimensions["C"].width = 40
out = run.path("review", "outgoing", "repair_rating_pack.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"{len(rows)} replies in {len(rows) // 2} pairs -> {out}; key -> review/repair_rating_key.csv")
