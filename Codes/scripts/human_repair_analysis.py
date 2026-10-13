"""Analyse the blinded human ratings of R2-T vs R5 repair replies (reviewer Q1).

    PYTHONPATH=src python scripts/human_repair_analysis.py
Reads review/human_returned/TLT_Reviewed_Another_Round/repair_rating_AGT-43.xlsx
and review/repair_rating_key.csv.  Writes results/tables/human_repair_*.csv.
"""
import csv
import json

import numpy as np
import openpyxl
from scipy.stats import binomtest, wilcoxon

from tlt.analysis.stats import add_holm, bootstrap_mean
from tlt.config import load_run

run = load_run("full")
B = int(run.design["bootstrap"]["resamples"])
SEED = run.seed
R = run.path("results")
T = R / "tables"

DIMS = ["Mistake_Identification", "Mistake_Location", "Revealing_of_the_Answer",
        "Providing_Guidance", "Math_Correctness", "Language_Naturalness"]
PQI_R_DIMS = DIMS[:5]

SCORE_MAP = {
    "Mistake_Identification": {"Yes": 1, "To some extent": 0.5, "No": 0},
    "Mistake_Location": {"Yes": 1, "To some extent": 0.5, "No": 0},
    "Revealing_of_the_Answer": {"No": 1, "Yes (correct)": 0, "Yes (incorrect)": 0},
    "Providing_Guidance": {"Yes": 1, "To some extent": 0.5, "No": 0},
    "Math_Correctness": {"Yes": 1, "No": 0},
    "Language_Naturalness": {"Natural": 1, "Understandable but unnatural": 0.5, "Hard to understand": 0},
}

key = {}
with open(run.path("review", "repair_rating_key.csv")) as f:
    for row in csv.DictReader(f):
        key[row["row_id"]] = row

wb = openpyxl.load_workbook(
    run.path("review", "human_returned", "TLT_Reviewed_Another_Round", "repair_rating_AGT-43.xlsx"))
ws = wb["Ratings"]
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
dim_cols = {d: headers.index(d) + 1 for d in DIMS}

# ---------------------------------------------------------------- 1. per-reply scores
import pandas as pd

rows = []
for r in range(2, ws.max_row + 1):
    row_id = ws.cell(r, 1).value
    k = key[row_id]
    rec = {"row_id": row_id, "pair": ws.cell(r, 2).value, "reply_label": ws.cell(r, 3).value,
           "language": k["language"], "arm": k["arm"], "item_id": k["item_id"], "model": k["model"]}
    for d in DIMS:
        raw = ws.cell(r, dim_cols[d]).value
        rec[d] = raw
        rec[f"s_{d}"] = SCORE_MAP[d][raw]
    rec["PQI_R"] = np.mean([rec[f"s_{d}"] for d in PQI_R_DIMS])
    rec["Language_Naturalness_score"] = rec["s_Language_Naturalness"]
    rows.append(rec)
df = pd.DataFrame(rows)

# ---------------------------------------------------------------- 2. arm-level summary
print("=== Per-arm means ===")
summary_rows = []
for arm in ("R2T", "R5"):
    sub = df[df.arm == arm]
    rec = {"arm": arm, "n": len(sub)}
    for d in DIMS:
        rec[d] = sub[f"s_{d}"].mean()
    rec["PQI_R"] = sub["PQI_R"].mean()
    summary_rows.append(rec)
    print(f"  {arm}: PQI-R={rec['PQI_R']:.3f}, n={rec['n']}")
t1 = pd.DataFrame(summary_rows)
t1.to_csv(T / "human_repair_arm_summary.csv", index=False)

# ---------------------------------------------------------------- 3. paired differences (within pair)
pairs = sorted(df.pair.unique())
pair_rows = []
for p in pairs:
    sub = df[df.pair == p]
    r2t = sub[sub.arm == "R2T"].iloc[0]
    r5 = sub[sub.arm == "R5"].iloc[0]
    rec = {"pair": p, "item_id": r2t["item_id"], "language": r2t["language"], "model": r2t["model"],
           "R2T_PQI_R": r2t["PQI_R"], "R5_PQI_R": r5["PQI_R"],
           "diff_PQI_R": r2t["PQI_R"] - r5["PQI_R"]}
    for d in DIMS:
        rec[f"diff_{d}"] = r2t[f"s_{d}"] - r5[f"s_{d}"]
    pair_rows.append(rec)
pdf = pd.DataFrame(pair_rows)
pdf.to_csv(T / "human_repair_paired.csv", index=False)

diffs = pdf["diff_PQI_R"].values
print(f"\n=== Paired PQI-R (R2T - R5), n={len(diffs)} pairs ===")
print(f"  Mean: {diffs.mean():+.3f}, Median: {np.median(diffs):+.3f}, SD: {diffs.std(ddof=1):.3f}")
print(f"  R2T>R5: {(diffs>0).sum()}, equal: {(diffs==0).sum()}, R5>R2T: {(diffs<0).sum()}")

nonzero = diffs[diffs != 0]
if len(nonzero) >= 5:
    w = wilcoxon(nonzero, alternative="two-sided")
    print(f"  Wilcoxon signed-rank (n={len(nonzero)}): W={w.statistic:.1f}, p={w.pvalue:.4f}")

bm = bootstrap_mean(pd.Series(diffs, index=range(len(diffs))), B, SEED)
print(f"  Bootstrap 95% CI: [{bm['ci_low']:+.3f}, {bm['ci_high']:+.3f}], p={bm['p_value']:.4f}")

# ---------------------------------------------------------------- 4. pair preferences (unblinded)
ps = wb["Pairs"]
pref_rows = []
for r in range(2, ps.max_row + 1):
    pair = ps.cell(r, 1).value
    choice = ps.cell(r, 3).value
    a_arm = key[pair + "A"]["arm"]
    b_arm = key[pair + "B"]["arm"]
    if choice == "A":
        winner = a_arm
    elif choice == "B":
        winner = b_arm
    else:
        winner = "tie"
    pref_rows.append({"pair": pair, "language": key[pair + "A"]["language"],
                       "model": key[pair + "A"]["model"], "chosen_label": choice, "winner": winner})
pf = pd.DataFrame(pref_rows)
pf.to_csv(T / "human_repair_preferences.csv", index=False)

pref_r2t = (pf.winner == "R2T").sum()
pref_r5 = (pf.winner == "R5").sum()
pref_tie = (pf.winner == "tie").sum()
decisive = pref_r2t + pref_r5
bt = binomtest(pref_r2t, decisive, 0.5, alternative="two-sided")
print(f"\n=== Pair preferences ===")
print(f"  R2-T preferred: {pref_r2t}/{len(pf)} ({pref_r2t/len(pf)*100:.0f}%)")
print(f"  R5 preferred:   {pref_r5}/{len(pf)} ({pref_r5/len(pf)*100:.0f}%)")
print(f"  Tie:            {pref_tie}/{len(pf)} ({pref_tie/len(pf)*100:.0f}%)")
print(f"  Binomial test ({decisive} decisive): p={bt.pvalue:.4f}")

for lang in ("HI", "BN"):
    sub = pf[pf.language == lang]
    r2t_w = (sub.winner == "R2T").sum()
    r5_w = (sub.winner == "R5").sum()
    tie_w = (sub.winner == "tie").sum()
    print(f"  {'Hindi' if lang=='HI' else 'Bengali'}: R2-T {r2t_w}, R5 {r5_w}, tie {tie_w}")

# ---------------------------------------------------------------- 5. by language
print("\n=== By language ===")
lang_rows = []
for lang in ("HI", "BN"):
    for arm in ("R2T", "R5"):
        sub = df[(df.language == lang) & (df.arm == arm)]
        rec = {"language": lang, "arm": arm, "n": len(sub), "PQI_R": sub.PQI_R.mean(),
               "Language_Naturalness": sub.Language_Naturalness_score.mean()}
        for d in PQI_R_DIMS:
            rec[d] = sub[f"s_{d}"].mean()
        lang_rows.append(rec)
        print(f"  {lang} {arm}: PQI-R={rec['PQI_R']:.3f}, Naturalness={rec['Language_Naturalness']:.3f}")
pd.DataFrame(lang_rows).to_csv(T / "human_repair_by_language.csv", index=False)

# ---------------------------------------------------------------- 6. by model
print("\n=== By model ===")
model_rows = []
for m in sorted(df.model.unique()):
    sub_pf = pf[pf.model == m]
    sub_pdf = pdf[pdf.model == m]
    r2t_w = (sub_pf.winner == "R2T").sum()
    r5_w = (sub_pf.winner == "R5").sum()
    mean_d = sub_pdf.diff_PQI_R.mean()
    model_rows.append({"model": m, "pairs": len(sub_pf), "prefer_R2T": r2t_w, "prefer_R5": r5_w,
                        "mean_PQI_R_diff": mean_d})
    print(f"  {m}: R2-T {r2t_w}, R5 {r5_w}; mean diff {mean_d:+.3f}")
pd.DataFrame(model_rows).to_csv(T / "human_repair_by_model.csv", index=False)

# ---------------------------------------------------------------- 7. agreement with automated judges
rep = pd.read_csv(R / "repair_judgements_aggregated.csv", low_memory=False)
rep = rep[(rep.variant == "original") & (rep.cue == "C0")]
auto = rep.set_index("target_id")
agree_rows = []
for _, row in df.iterrows():
    tid = row["row_id"]  # we need the target_id from the key
    k = key[tid]
    target_id = k["target_id"]
    if target_id in auto.index:
        a = auto.loc[target_id]
        agree_rows.append({
            "row_id": tid, "arm": row["arm"], "language": row["language"],
            "human_PQI_R": row["PQI_R"],
            "auto_PQI_R": a["PQI_R"] if "PQI_R" in a.index else np.nan,
        })
adf = pd.DataFrame(agree_rows)
if len(adf) > 0 and "auto_PQI_R" in adf and adf.auto_PQI_R.notna().sum() > 0:
    corr = adf[["human_PQI_R", "auto_PQI_R"]].corr().iloc[0, 1]
    mad = (adf.human_PQI_R - adf.auto_PQI_R).abs().mean()
    print(f"\n=== Human vs automated PQI-R ===")
    print(f"  Correlation: {corr:.3f}")
    print(f"  Mean absolute difference: {mad:.3f}")
    print(f"  Human mean: {adf.human_PQI_R.mean():.3f}, Auto mean: {adf.auto_PQI_R.mean():.3f}")
    for arm in ("R2T", "R5"):
        s = adf[adf.arm == arm]
        print(f"  {arm}: human={s.human_PQI_R.mean():.3f}, auto={s.auto_PQI_R.mean():.3f}")
    adf.to_csv(T / "human_repair_vs_auto.csv", index=False)

# ---------------------------------------------------------------- 8. combined results table
results = {
    "n_replies": len(df),
    "n_pairs": len(pairs),
    "R2T_PQI_R": df[df.arm == "R2T"].PQI_R.mean(),
    "R5_PQI_R": df[df.arm == "R5"].PQI_R.mean(),
    "paired_diff_mean": diffs.mean(),
    "paired_diff_ci_low": bm["ci_low"],
    "paired_diff_ci_high": bm["ci_high"],
    "paired_diff_p": bm["p_value"],
    "wilcoxon_p": w.pvalue if len(nonzero) >= 5 else np.nan,
    "prefer_R2T": int(pref_r2t),
    "prefer_R5": int(pref_r5),
    "prefer_tie": int(pref_tie),
    "binomial_p": bt.pvalue,
}
pd.DataFrame([results]).to_csv(T / "human_repair_summary.csv", index=False)
print("\n=== Summary ===")
for k, v in results.items():
    print(f"  {k}: {v}")
print("\nDone. Tables written to results/tables/human_repair_*.csv")
